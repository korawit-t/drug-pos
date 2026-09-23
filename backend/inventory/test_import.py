import io

from django.test import Client, TestCase
from openpyxl import Workbook, load_workbook

from accounts.models import User

from .importer import COLUMNS, import_rows, read_rows
from .models import Lot, Product, ProductUnit, Purchase

HEADERS = [names[0] for names in COLUMNS.values()]


def row(**values) -> list:
    """A spreadsheet row in template column order, written as field=value."""
    cells = dict.fromkeys(COLUMNS, "")
    cells.update(values)
    return [cells[field] for field in COLUMNS]


def sheet(rows, headers=HEADERS) -> bytes:
    workbook = Workbook()
    workbook.active.append(headers)
    for line in rows:
        workbook.active.append(line)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


AMOX_STRIP = dict(
    trade_name="Amoxicillin", generic_name="Amoxicillin trihydrate", strength="500 mg", dosage_form="แคปซูล",
    category="ยาอันตราย", base_unit="แคปซูล", unit_name="แผง", factor=10, barcode="200042", is_default="ใช่",
    price1=40, price2=38, storage_location="ชั้น B2",
)
AMOX_BOX = dict(
    trade_name="Amoxicillin", strength="500 mg", base_unit="แคปซูล", unit_name="กล่อง", factor=100,
    barcode="200059", price1=380,
)


def run(rows, *, with_stock=False, commit=True, user=None):
    return import_rows(read_rows(sheet(rows)), with_stock=with_stock, user=user, commit=commit, source="test.xlsx")


class ImportProductsTests(TestCase):
    def test_rows_of_one_drug_become_one_product_with_its_units(self):
        result = run([row(**AMOX_STRIP), row(**AMOX_BOX)])
        self.assertTrue(result.ok, result.errors)
        self.assertEqual((result.products_created, result.units_created), (1, 2))
        product = Product.objects.get()
        self.assertEqual(product.generic_name, "Amoxicillin trihydrate")
        self.assertEqual(product.category, Product.Category.DANGEROUS)
        strip = product.units.get(name="แผง")
        self.assertEqual((strip.factor, strip.barcode, strip.price1, strip.price2), (10, "200042", 40, 38))
        self.assertTrue(strip.is_default)
        self.assertFalse(product.units.get(name="กล่อง").is_default)

    def test_a_drug_with_two_unit_rows_counts_as_one_new_drug(self):
        result = run([row(**AMOX_STRIP), row(**AMOX_BOX)])
        self.assertEqual((result.products_created, result.products_updated), (1, 0))

    def test_checking_the_file_writes_nothing(self):
        result = run([row(**AMOX_STRIP)], commit=False)
        self.assertTrue(result.ok)
        self.assertEqual(result.products_created, 1)
        self.assertFalse(result.committed)
        self.assertFalse(Product.objects.exists())

    def test_importing_again_updates_instead_of_duplicating(self):
        run([row(**AMOX_STRIP)])
        result = run([row(**{**AMOX_STRIP, "price1": 45, "generic_name": ""})])
        self.assertEqual((result.products_created, result.units_created), (0, 0))
        self.assertEqual(Product.objects.count(), 1)
        self.assertEqual(ProductUnit.objects.get().price1, 45)
        # An empty cell leaves what was there alone.
        self.assertEqual(Product.objects.get().generic_name, "Amoxicillin trihydrate")

    def test_barcode_belonging_to_another_drug_is_refused(self):
        run([row(**AMOX_STRIP)])
        result = run([row(**{**AMOX_STRIP, "trade_name": "Paracetamol", "generic_name": "Paracetamol"})])
        self.assertFalse(result.ok)
        self.assertIn("บาร์โค้ด", result.errors[0].message)
        self.assertEqual(Product.objects.count(), 1)

    def test_the_same_unit_twice_in_one_file_is_refused(self):
        result = run([row(**AMOX_STRIP), row(**{**AMOX_STRIP, "barcode": "200999", "price1": 44})])
        self.assertFalse(result.ok)
        self.assertIn("ซ้ำกับแถวที่ 2", result.errors[0].message)
        self.assertFalse(Product.objects.exists())

    def test_one_bad_row_stops_the_whole_file(self):
        result = run([row(**AMOX_STRIP), row(**{**AMOX_BOX, "factor": "สิบ"})])
        self.assertFalse(result.ok)
        self.assertEqual(result.errors[0].row, 3)
        self.assertFalse(Product.objects.exists())

    def test_a_new_drug_needs_a_category_and_a_known_one(self):
        self.assertIn("ประเภท", run([row(**{**AMOX_STRIP, "category": ""})]).errors[0].message)
        self.assertIn("ไม่รู้จักประเภท", run([row(**{**AMOX_STRIP, "category": "ยาแรง"})]).errors[0].message)

    def test_price_one_is_required(self):
        self.assertIn("ราคา 1", run([row(**{**AMOX_STRIP, "price1": ""})]).errors[0].message)


class ImportStockTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("owner", password="x", is_staff=True)

    def test_opening_stock_is_posted_as_one_receipt(self):
        result = run(
            [row(**{**AMOX_STRIP, "stock_qty": 5, "lot_no": "AX118", "expiry": "03/2571", "unit_cost": 230})],
            with_stock=True, user=self.user,
        )
        self.assertTrue(result.ok, result.errors)
        self.assertEqual((result.stock_lines, result.stock_base_qty), (1, 50))
        purchase = Purchase.objects.get()
        self.assertTrue(purchase.is_opening_balance)
        self.assertEqual(purchase.status, Purchase.Status.POSTED)
        lot = Lot.objects.get()
        self.assertEqual((lot.lot_no, str(lot.expiry_date), lot.qty_on_hand), ("AX118", "2028-03-31", 50))
        self.assertEqual(lot.movements.get().qty, 50)

    def test_stock_columns_are_ignored_when_not_asked_for(self):
        result = run([row(**{**AMOX_STRIP, "stock_qty": 5, "lot_no": "AX118", "expiry": "03/2571"})])
        self.assertEqual(result.stock_lines, 0)
        self.assertFalse(Lot.objects.exists())
        self.assertTrue(Product.objects.exists())

    def test_stock_needs_a_lot_and_a_readable_expiry(self):
        result = run([row(**{**AMOX_STRIP, "stock_qty": 5, "lot_no": "AX118", "expiry": "เดือนหน้า"})], with_stock=True)
        self.assertFalse(result.ok)
        self.assertIn("lot", result.errors[0].message)
        self.assertFalse(Product.objects.exists())

    def test_expired_stock_is_taken_with_a_warning(self):
        result = run(
            [row(**{**AMOX_STRIP, "stock_qty": 1, "lot_no": "OLD", "expiry": "01/2560"})],
            with_stock=True, user=self.user,
        )
        self.assertTrue(result.ok, result.errors)
        self.assertIn("หมดอายุแล้ว", result.warnings[0].message)
        self.assertEqual(Lot.objects.get().qty_on_hand, 10)


class ReadFileTests(TestCase):
    def test_csv_saved_by_excel_in_thai_windows(self):
        text = "ชื่อการค้า,ประเภท,หน่วยเล็กสุด,หน่วยขาย,ตัวคูณ,ราคา\nพารา,ยาสามัญประจำบ้าน,เม็ด,แผง,10,15\n"
        rows = read_rows(text.encode("cp874"), "stock.csv")
        self.assertEqual(rows[0]["trade_name"], "พารา")
        self.assertEqual(rows[0]["price1"], "15")
        result = import_rows(rows, with_stock=False, commit=True)
        self.assertTrue(result.ok, result.errors)
        self.assertEqual(Product.objects.get().base_unit, "เม็ด")

    def test_a_file_without_the_needed_headers_says_so(self):
        with self.assertRaises(Exception) as ctx:
            read_rows(sheet([["a", "b"]], headers=["สิ่งของ", "จำนวน"]))
        self.assertIn("หัวตาราง", str(ctx.exception))


class ImportApiTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user("owner", password="x", is_staff=True)
        self.client = Client(enforce_csrf_checks=True)
        self.client.force_login(self.owner)
        self.client.get("/api/auth/csrf")

    def upload(self, data, **fields):
        return self.client.post(
            "/api/import/products",
            {"file": io.BytesIO(data), **fields},
            HTTP_X_CSRFTOKEN=self.client.cookies["csrftoken"].value,
        )

    def test_template_has_a_help_sheet_and_can_be_imported_back(self):
        response = self.client.get("/api/import/template")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["ยา", "คำอธิบาย"])
        result = import_rows(read_rows(response.content), with_stock=True, user=self.owner, commit=True)
        self.assertTrue(result.ok, result.errors)
        self.assertEqual(Product.objects.count(), 2)  # the template's example drugs

    def test_check_then_import(self):
        data = sheet([row(**AMOX_STRIP)])
        checked = self.upload(data, with_stock="false", commit="false").json()
        self.assertEqual(checked["products_created"], 1)
        self.assertFalse(checked["committed"])
        self.assertFalse(Product.objects.exists())

        imported = self.upload(data, commit="true").json()
        self.assertTrue(imported["committed"])
        self.assertEqual(Product.objects.count(), 1)

    def test_errors_come_back_with_row_numbers(self):
        response = self.upload(sheet([row(**{**AMOX_STRIP, "price1": "ฟรี"})]), commit="true")
        body = response.json()
        self.assertEqual(body["error_count"], 1)
        self.assertEqual(body["errors"][0]["row"], 2)
        self.assertFalse(body["committed"])

    def test_staff_only(self):
        self.client.force_login(User.objects.create_user("cashier", password="x"))
        self.client.get("/api/auth/csrf")
        self.assertEqual(self.upload(sheet([row(**AMOX_STRIP)]), commit="true").status_code, 403)
